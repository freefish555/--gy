#!/usr/bin/env python3
"""Document 1: System Function & Design Architecture"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

SCREENSHOTS_DIR = "/home/user/webapp/docs/screenshots"
OUTPUT_DIR = "/home/user/webapp/docs"

def set_cell_background(cell, color):
    """Set table cell background color"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def set_table_borders(table):
    """Add borders to table"""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_page_break(doc):
    para = doc.add_paragraph()
    run = para.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)
    return para

def heading(doc, text, level=1, color='1F3864'):
    p = doc.add_heading(text, level=level)
    if color:
        for run in p.runs:
            run.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_para(doc, text, bold=False, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    run.font.name = '宋体'
    run.font.size = Pt(11)
    run.font.bold = bold
    return p

def add_screenshot(doc, filename, caption="", width=15):
    path = os.path.join(SCREENSHOTS_DIR, filename)
    if os.path.exists(path):
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(path, width=Cm(width))
            if caption:
                cap_p = doc.add_paragraph(caption)
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cap_p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(100, 100, 100)
        except Exception as e:
            doc.add_paragraph(f"[截图: {caption}]")
    else:
        doc.add_paragraph(f"[截图文件不存在: {filename}]")

def create_doc1():
    doc = Document()
    
    # Page setup
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    
    # Default style
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)
    
    # ===== COVER PAGE =====
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run('等保项目管理及归档管理系统')
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 56, 100)
    title_run.font.name = '黑体'
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run('系统功能及开发设计架构说明')
    sub_run.font.size = Pt(18)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(31, 56, 100)
    sub_run.font.name = '黑体'
    
    doc.add_paragraph()
    
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info_table)
    info_data = [
        ('编制单位', '南京国云电力有限公司'),
        ('系统版本', 'v2.0'),
        ('文档版本', 'v2.0'),
        ('编制日期', '2026年3月'),
        ('密级', '内部资料'),
    ]
    for i, (k, v) in enumerate(info_data):
        row = info_table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        set_cell_background(row.cells[0], 'D6E4F7')
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
    
    add_page_break(doc)
    
    # ===== TABLE OF CONTENTS =====
    heading(doc, '目录', 1)
    toc_items = [
        ('1', '系统概述', '3'),
        ('1.1', '项目背景', '3'),
        ('1.2', '系统定位与目标', '3'),
        ('1.3', '系统特点', '3'),
        ('2', '系统架构设计', '4'),
        ('2.1', '整体技术架构', '4'),
        ('2.2', '前端技术栈', '4'),
        ('2.3', '后端技术栈', '4'),
        ('2.4', '数据安全设计', '5'),
        ('2.5', '接口设计规范', '5'),
        ('3', '系统功能模块说明', '6'),
        ('3.1', '用户登录与认证', '6'),
        ('3.2', '项目管理模块', '7'),
        ('3.3', '归档材料制作模块', '9'),
        ('3.4', '基础数据管理', '10'),
        ('3.5', '系统管理', '12'),
        ('3.6', '日志管理', '13'),
        ('4', '数据库设计说明', '14'),
        ('4.1', '数据库概览', '14'),
        ('4.2', '核心数据表详细说明', '14'),
        ('5', '数据加密字段说明', '20'),
        ('6', '系统部署结构', '21'),
    ]
    for num, title, page in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5 if '.' in num and num.count('.') == 1 else (1.0 if num.count('.') > 1 else 0))
        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Cm(14), leader=1)
        run1 = p.add_run(f'{num}  {title}')
        run1.font.size = Pt(11)
        run2 = p.add_run(f'\t{page}')
        run2.font.size = Pt(11)
    
    add_page_break(doc)
    
    # ===== CHAPTER 1: OVERVIEW =====
    heading(doc, '1  系统概述', 1)
    
    heading(doc, '1.1  项目背景', 2)
    add_para(doc, '随着网络安全等级保护（等保）工作的深入推进，测评机构需要对大量项目进行全流程管理，包括项目立项、人员配置、被测系统登记、测评过程管控以及最终归档材料的制作与管理。传统的手工管理方式效率低下、信息分散，难以满足合规性要求和管理需要。')
    add_para(doc, '南京国云电力有限公司开发了本系统，旨在通过信息化手段将等保项目管理全流程数字化，实现项目信息的集中管理、归档材料的自动生成，以及安全合规的数据保护。')
    
    heading(doc, '1.2  系统定位与目标', 2)
    add_para(doc, '本系统（等保项目管理及归档管理系统 v2.0）是面向网络安全等级保护测评机构的内部项目管理系统，主要目标包括：')
    items = [
        '实现等保测评项目的全生命周期管理（立项、配置人员、测评、归档）；',
        '自动生成归档材料Word文档，支持多种模板（保密协议、任务书、计划书、报告等）；',
        '对涉密信息（客户名称、联系方式、合同金额等）进行SM4加密存储；',
        '支持TOTP双因素认证，保障系统访问安全；',
        '提供统计分析功能，支持按项目类型、行业、人员维度进行数据统计；',
        '完整的操作日志审计，满足合规要求。',
    ]
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    heading(doc, '1.3  系统特点', 2)
    features = [
        ('安全性高', 'SM4国密算法加密存储敏感字段；JWT + TOTP双因素认证；操作日志全程审计；'),
        ('自动化程度高', '基于Word模板（python-docx/Apache POI）自动填充项目信息，生成标准化归档材料包；'),
        ('前后端分离', '采用Vue 3 + Element Plus前端与Spring Boot后端分离架构，接口标准化（RESTful API）；'),
        ('数据完整性', '关键数据写入时计算HMAC摘要，防止数据被篡改；'),
        ('灵活扩展', '模板管理支持上传自定义模板，字典管理支持动态配置分类项目；'),
    ]
    feat_table = doc.add_table(rows=len(features)+1, cols=2)
    set_table_borders(feat_table)
    feat_table.rows[0].cells[0].text = '特点'
    feat_table.rows[0].cells[1].text = '说明'
    set_cell_background(feat_table.rows[0].cells[0], '1F3864')
    set_cell_background(feat_table.rows[0].cells[1], '1F3864')
    for c in feat_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(features):
        row = feat_table.rows[i+1]
        row.cells[0].text = k
        row.cells[1].text = v
        if i % 2 == 0:
            set_cell_background(row.cells[0], 'EBF3FB')
            set_cell_background(row.cells[1], 'EBF3FB')
    
    add_page_break(doc)
    
    # ===== CHAPTER 2: ARCHITECTURE =====
    heading(doc, '2  系统架构设计', 1)
    
    heading(doc, '2.1  整体技术架构', 2)
    add_para(doc, '系统采用经典的前后端分离三层架构，分为展示层、业务逻辑层和数据持久层：')
    
    arch_table = doc.add_table(rows=4, cols=3)
    set_table_borders(arch_table)
    arch_headers = ['层次', '组件', '说明']
    for j, h in enumerate(arch_headers):
        arch_table.rows[0].cells[j].text = h
        set_cell_background(arch_table.rows[0].cells[j], '1F3864')
        for p in arch_table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    arch_data = [
        ('展示层（前端）', 'Vue 3 + Vite + Element Plus', 'SPA单页应用，Nginx静态服务'),
        ('业务逻辑层（后端）', 'Spring Boot 3 + MyBatis-Plus', 'RESTful API，JWT认证，SM4加密'),
        ('数据持久层', 'MySQL 8.0', '关系型数据库，存储所有业务数据'),
    ]
    for i, (a, b, c) in enumerate(arch_data):
        arch_table.rows[i+1].cells[0].text = a
        arch_table.rows[i+1].cells[1].text = b
        arch_table.rows[i+1].cells[2].text = c
    
    heading(doc, '2.2  前端技术栈', 2)
    fe_table = doc.add_table(rows=8, cols=3)
    set_table_borders(fe_table)
    fe_headers = ['技术/框架', '版本', '用途']
    for j, h in enumerate(fe_headers):
        fe_table.rows[0].cells[j].text = h
        set_cell_background(fe_table.rows[0].cells[j], '2E75B6')
        for p in fe_table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    fe_data = [
        ('Vue 3', '3.x', '前端框架，Composition API'),
        ('Vite', '5.x', '构建工具，开发热更新'),
        ('Element Plus', '2.x', 'UI组件库'),
        ('Pinia', '2.x', '状态管理'),
        ('Vue Router', '4.x', '前端路由'),
        ('Axios', '1.x', 'HTTP请求库'),
        ('ECharts', '5.x', '数据可视化图表'),
    ]
    for i, row_data in enumerate(fe_data):
        for j, val in enumerate(row_data):
            fe_table.rows[i+1].cells[j].text = val
    
    heading(doc, '2.3  后端技术栈', 2)
    be_table = doc.add_table(rows=9, cols=3)
    set_table_borders(be_table)
    be_headers = ['技术/框架', '版本', '用途']
    for j, h in enumerate(be_headers):
        be_table.rows[0].cells[j].text = h
        set_cell_background(be_table.rows[0].cells[j], '2E75B6')
        for p in be_table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    be_data = [
        ('Spring Boot 3', '3.x', '核心框架，自动配置，依赖注入'),
        ('MyBatis-Plus', '3.5.x', 'ORM框架，简化数据库操作'),
        ('MySQL 8', '8.0.x', '关系型数据库'),
        ('Java', '17 (LTS)', '运行环境'),
        ('JWT (jjwt)', '0.12.x', 'Token认证，无状态会话'),
        ('Bouncy Castle', '1.77', 'SM4国密算法加密'),
        ('Apache POI', '5.x', 'Word文档生成（归档材料）'),
        ('Lombok', '最新', '代码简化注解'),
    ]
    for i, row_data in enumerate(be_data):
        for j, val in enumerate(row_data):
            be_table.rows[i+1].cells[j].text = val
    
    heading(doc, '2.4  数据安全设计', 2)
    add_para(doc, '系统在数据安全方面采用多层保护机制：')
    security_items = [
        '【传输层安全】前后端之间采用HTTPS加密传输，防止数据窃听；',
        '【存储层加密】涉及个人隐私和商业敏感的字段采用SM4（国密）对称加密算法存储，加密密钥由环境变量 APP_ROOT_KEY 配置；',
        '【数据完整性】关键业务数据写入时计算HMAC-SHA256摘要（dataHmac字段），防止数据被篡改；',
        '【认证安全】采用JWT + TOTP（基于时间的一次性密码）双因素认证，令牌有效期8小时；',
        '【账户保护】连续登录失败5次后账户锁定，防止暴力破解；',
        '【权限控制】基于RBAC角色权限控制，系统内置管理员、普通用户等角色，细粒度控制接口权限；',
    ]
    for item in security_items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    heading(doc, '2.5  接口设计规范', 2)
    add_para(doc, '系统采用RESTful API设计规范，接口基础路径为 /api。')
    
    api_table = doc.add_table(rows=7, cols=3)
    set_table_borders(api_table)
    api_headers = ['接口路径', 'HTTP方法', '说明']
    for j, h in enumerate(api_headers):
        api_table.rows[0].cells[j].text = h
        set_cell_background(api_table.rows[0].cells[j], '2E75B6')
        for p in api_table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    api_data = [
        ('/api/auth/login', 'POST', '用户登录，返回JWT令牌'),
        ('/api/project/page', 'GET', '项目分页查询，支持多条件筛选'),
        ('/api/project/{id}', 'GET', '获取项目详情'),
        ('/api/project', 'POST/PUT', '创建/更新项目'),
        ('/api/archive/generate', 'POST', '生成归档材料压缩包'),
        ('/api/archive/templates', 'GET', '获取可用归档模板列表'),
    ]
    for i, (a, b, c) in enumerate(api_data):
        api_table.rows[i+1].cells[0].text = a
        api_table.rows[i+1].cells[1].text = b
        api_table.rows[i+1].cells[2].text = c
    
    add_para(doc, '\n统一响应格式：')
    code_para = doc.add_paragraph()
    code_run = code_para.add_run('{ "code": 200, "message": "success", "data": { ... } }')
    code_run.font.name = 'Courier New'
    code_run.font.size = Pt(10)
    code_para.paragraph_format.left_indent = Cm(1)
    
    add_page_break(doc)
    
    # ===== CHAPTER 3: FUNCTIONAL MODULES =====
    heading(doc, '3  系统功能模块说明', 1)
    
    heading(doc, '3.1  用户登录与认证', 2)
    add_para(doc, '用户登录界面提供账户/密码输入，支持图形验证码（防机器人），可选择开启TOTP双因素认证。')
    
    features_auth = [
        ('标准登录', '用户名+密码+图形验证码登录，登录失败超5次锁定账户'),
        ('TOTP认证', '支持开启基于时间的一次性密码（兼容Google Authenticator等App）'),
        ('JWT令牌', '登录成功返回JWT令牌，有效期8小时，接口认证无状态'),
        ('登录日志', '记录每次登录的IP地址、时间、成功/失败状态'),
        ('密码修改', '首次登录强制修改密码，支持随时修改'),
    ]
    auth_table = doc.add_table(rows=len(features_auth)+1, cols=2)
    set_table_borders(auth_table)
    auth_table.rows[0].cells[0].text = '功能点'
    auth_table.rows[0].cells[1].text = '说明'
    set_cell_background(auth_table.rows[0].cells[0], '1F3864')
    set_cell_background(auth_table.rows[0].cells[1], '1F3864')
    for c in auth_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(features_auth):
        auth_table.rows[i+1].cells[0].text = k
        auth_table.rows[i+1].cells[1].text = v
    
    doc.add_paragraph()
    add_screenshot(doc, '01_login.png', '图3-1 用户登录界面', 14)
    
    heading(doc, '3.2  项目管理模块', 2)
    add_para(doc, '项目管理模块是系统的核心功能，提供等保测评项目的全生命周期管理。')
    
    heading(doc, '3.2.1  项目总览（列表）', 3)
    add_para(doc, '以表格形式展示所有项目，支持快速搜索和高级搜索过滤，支持按多个维度进行筛选。')
    
    search_fields = [
        ('普通搜索', '项目编号、项目名称、合同日期范围'),
        ('高级搜索', '备案编号、委托单位、项目负责人、项目经理、登记测评师、项目类型、所属行业、年份'),
        ('项目组成员', '按登记测评师姓名模糊搜索，显示该人员参与的所有项目'),
        ('排序', '支持按创建时间正/倒序排列'),
        ('分页', '支持自定义每页显示条数（10/20/50/100）'),
    ]
    sf_table = doc.add_table(rows=len(search_fields)+1, cols=2)
    set_table_borders(sf_table)
    sf_table.rows[0].cells[0].text = '搜索类型'
    sf_table.rows[0].cells[1].text = '支持字段'
    set_cell_background(sf_table.rows[0].cells[0], '2E75B6')
    set_cell_background(sf_table.rows[0].cells[1], '2E75B6')
    for c in sf_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(search_fields):
        sf_table.rows[i+1].cells[0].text = k
        sf_table.rows[i+1].cells[1].text = v
    
    doc.add_paragraph()
    add_screenshot(doc, '03_project_list.png', '图3-2 项目总览列表', 14)
    add_screenshot(doc, '04_project_advanced_search.png', '图3-3 高级搜索面板', 14)
    
    heading(doc, '3.2.2  项目详情', 3)
    add_para(doc, '项目详情页面以分组形式展示项目的完整信息，分为以下信息区块：')
    detail_groups = [
        '基本信息：项目编号、项目名称、委托单位、客户地址、合同日期、合同金额、项目类型、所属行业；',
        '被测系统列表：系统序号、系统名称、安全等级（二级/三级）、测评指标、备案编号；',
        '项目人员：按角色分类展示（项目经理、项目负责人、登记测评师、实际参与成员等12种角色）；',
        '归档状态：纸质归档状态、电子归档状态；',
        '进度追踪：准备阶段、计划阶段、现场测评、报告阶段的完成日期；',
        '其他信息：任务书签约时间、所属年份、业务人员、项目地区、报告邮寄日期、报告邮寄编号、备注。',
    ]
    for item in detail_groups:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    doc.add_paragraph()
    add_screenshot(doc, '05_project_detail.png', '图3-4 项目详情页面', 14)
    
    heading(doc, '3.2.3  新建/编辑项目', 3)
    add_para(doc, '提供完整的项目创建和编辑表单，支持动态添加被测系统（可添加多个），人员配置支持按角色进行多对多关联。')
    
    heading(doc, '3.2.4  项目导入/导出', 3)
    add_para(doc, '支持通过Excel模板批量导入项目数据，同时支持将查询结果导出为Excel格式。')
    
    heading(doc, '3.2.5  项目统计', 3)
    add_para(doc, '以图表形式展示项目统计数据，包括：')
    stats_items = [
        '按项目经理统计：各项目经理名下的项目数量（柱状图）；',
        '按项目类型统计：各类型项目的数量分布（饼图）；',
        '按所属行业统计：各行业项目的数量分布（饼图）；',
        '按等级统计：二级/三级系统数量统计；',
    ]
    for item in stats_items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    add_screenshot(doc, '07_project_stats.png', '图3-5 项目统计图表', 14)
    
    heading(doc, '3.3  归档材料制作模块', 2)
    add_para(doc, '归档材料制作模块是系统的核心自动化功能，能够根据项目信息自动填充Word模板，批量生成标准化归档材料压缩包。')
    
    heading(doc, '3.3.1  功能说明', 3)
    archive_features = [
        ('项目选择', '通过关键词搜索（支持项目编号或名称模糊匹配）快速定位目标项目'),
        ('模板选择', '展示系统内所有可用归档模板，可勾选需要生成的材料'),
        ('自动生成', '系统根据项目信息自动替换Word模板中的占位符，生成对应文档'),
        ('批量下载', '将所有生成的文档打包为ZIP压缩包，以"项目编号-归档材料.zip"命名下载'),
        ('保密协议特殊处理', '5号文件（保密协议等）中含下划线格式的字段，系统特别处理以保留下划线格式'),
    ]
    af_table = doc.add_table(rows=len(archive_features)+1, cols=2)
    set_table_borders(af_table)
    af_table.rows[0].cells[0].text = '功能点'
    af_table.rows[0].cells[1].text = '说明'
    set_cell_background(af_table.rows[0].cells[0], '1F3864')
    set_cell_background(af_table.rows[0].cells[1], '1F3864')
    for c in af_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(archive_features):
        af_table.rows[i+1].cells[0].text = k
        af_table.rows[i+1].cells[1].text = v
    
    heading(doc, '3.3.2  模板占位符规范', 3)
    add_para(doc, '归档模板中使用 {{变量名}} 或 ${变量名} 格式的占位符，系统在生成时自动替换为实际数据。常用占位符说明：')
    
    placeholder_data = [
        ('{{xmbh}}', '项目编号'),
        ('{{xmmc}}', '项目名称'),
        ('{{khmc}}', '委托单位名称'),
        ('{{xtmc}}', '被测系统名称（多个系统用顿号分隔）'),
        ('{{xmfzr}}', '项目负责人姓名'),
        ('{{xmjl}}', '项目经理姓名'),
        ('{{djcps}}', '登记测评师姓名'),
        ('{{bah1}}, {{bah2}}...', '各被测系统备案编号（按序号索引）'),
        ('{{sys1_name}}...', '各被测系统名称（按序号索引）'),
        ('{{contractDate}}', '合同签署日期'),
    ]
    ph_table = doc.add_table(rows=len(placeholder_data)+1, cols=2)
    set_table_borders(ph_table)
    ph_table.rows[0].cells[0].text = '占位符'
    ph_table.rows[0].cells[1].text = '对应字段'
    set_cell_background(ph_table.rows[0].cells[0], '2E75B6')
    set_cell_background(ph_table.rows[0].cells[1], '2E75B6')
    for c in ph_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(placeholder_data):
        ph_table.rows[i+1].cells[0].text = k
        ph_table.rows[i+1].cells[1].text = v
    
    doc.add_paragraph()
    add_screenshot(doc, '08_archive_page.png', '图3-6 归档材料制作界面', 14)
    
    add_page_break(doc)
    
    heading(doc, '3.4  基础数据管理', 2)
    
    heading(doc, '3.4.1  人员管理', 3)
    add_para(doc, '管理等保测评人员（t_staff）档案，包括：')
    staff_fields = [
        '人员编号（唯一标识）、姓名（SM4加密）；',
        '部门、职位、资质等级（高级/中级/初级）；',
        '证书编号（SM4加密）、证书有效期；',
        '联系电话（SM4加密）、电子邮箱（SM4加密）；',
        '人员状态（在职/离职）；',
        '支持批量导入（Excel格式）和导出功能。',
    ]
    for item in staff_fields:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    add_screenshot(doc, '09_staff_manage.png', '图3-7 人员管理界面', 14)
    
    heading(doc, '3.4.2  测评设备管理', 3)
    add_para(doc, '管理硬件测评设备（t_eval_device）档案，包括：')
    device_info = [
        '设备编号、设备名称、设备型号；',
        '设备类型（测评专用机/扫描设备）；',
        '归属人员（关联t_staff人员）；',
        '设备状态（正常/停用）；',
        '归档材料生成时，工具清单中会自动显示项目对应的扫描设备信息。',
    ]
    for item in device_info:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    add_screenshot(doc, '10_device_manage.png', '图3-8 测评设备管理界面', 14)
    
    heading(doc, '3.4.3  归档模板管理', 3)
    add_para(doc, '管理归档材料Word模板（t_archive_template），支持：')
    template_info = [
        '上传自定义Word模板（.doc/.docx格式）；',
        '模板分类管理（按项目类型和系统等级适用）；',
        '模板版本控制；',
        '模板启用/停用管理；',
        '模板内容预览（通过下载查看）；',
        '模板替换测试（上传文档并用指定项目数据测试替换效果）。',
    ]
    for item in template_info:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    add_screenshot(doc, '11_template_manage.png', '图3-9 归档模板管理界面', 14)
    
    heading(doc, '3.4.4  字典管理', 3)
    add_para(doc, '系统通过字典（t_dict / t_dict_item）管理各类下拉选项，当前字典类型包括：')
    dict_types = [
        ('project_type', '项目类型（如：电力行业、金融行业等）'),
        ('industry', '所属行业分类'),
        ('project_region', '项目所属地区'),
        ('role_level', '测评师资质等级'),
        ('phase_status', '项目阶段状态'),
    ]
    dt_table = doc.add_table(rows=len(dict_types)+1, cols=2)
    set_table_borders(dt_table)
    dt_table.rows[0].cells[0].text = '字典编码'
    dt_table.rows[0].cells[1].text = '说明'
    set_cell_background(dt_table.rows[0].cells[0], '2E75B6')
    set_cell_background(dt_table.rows[0].cells[1], '2E75B6')
    for c in dt_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(dict_types):
        dt_table.rows[i+1].cells[0].text = k
        dt_table.rows[i+1].cells[1].text = v
    
    add_screenshot(doc, '13_dict_manage.png', '图3-10 字典管理界面', 14)
    
    heading(doc, '3.5  系统管理', 2)
    
    heading(doc, '3.5.1  用户管理', 3)
    add_para(doc, '管理系统登录用户（t_user），包括：')
    user_mgmt = [
        '创建/编辑/停用系统用户，关联角色；',
        '用户真实姓名、联系方式均SM4加密存储；',
        '支持重置用户密码；',
        '支持启用/禁用TOTP双因素认证；',
        '查看用户最近登录时间和登录IP。',
    ]
    for item in user_mgmt:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    add_screenshot(doc, '12_user_manage.png', '图3-11 用户管理界面', 14)
    
    heading(doc, '3.5.2  系统配置', 3)
    add_para(doc, '系统配置（t_sys_config）提供可在界面上修改的系统参数，包括：')
    config_items = [
        '系统名称（SYS_NAME）：显示在系统标题栏的名称；',
        '公司LOGO（SYS_LOGO）：上传自定义公司Logo图片；',
        '其他业务参数可通过配置项扩展。',
    ]
    for item in config_items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item).font.size = Pt(11)
    
    add_screenshot(doc, '14_system_config.png', '图3-12 系统配置界面', 14)
    
    heading(doc, '3.6  日志管理', 2)
    
    heading(doc, '3.6.1  登录日志', 3)
    add_para(doc, '记录所有用户的登录行为（t_login_log），包括登录用户名、真实姓名、登录IP、登录时间、登录状态（成功/失败）及失败原因，支持按用户名、状态、时间范围查询。')
    add_screenshot(doc, '15_login_log.png', '图3-13 登录日志界面', 14)
    
    heading(doc, '3.6.2  操作日志', 3)
    add_para(doc, '记录所有用户的关键操作行为（t_operation_log），包括操作用户、操作类型（CREATE/UPDATE/DELETE/EXPORT等）、操作对象（项目/人员/设备等）、操作描述、操作IP和操作时间，支持按用户、操作类型、时间范围查询。')
    add_screenshot(doc, '16_operation_log.png', '图3-14 操作日志界面', 14)
    
    add_page_break(doc)
    
    # ===== CHAPTER 4: DATABASE =====
    heading(doc, '4  数据库设计说明', 1)
    
    heading(doc, '4.1  数据库概览', 2)
    add_para(doc, '系统使用 MySQL 8.0 数据库，数据库名称为 djbh_system，共包含18张数据表：')
    
    db_overview = [
        ('t_project', '项目主表', '存储项目的核心信息'),
        ('t_project_system', '被测系统表', '每个项目可包含多个被测系统'),
        ('t_project_member', '项目人员关联表', '项目与人员的多对多关联，含角色类型'),
        ('t_staff', '人员清单表', '测评人员基础档案'),
        ('t_user', '系统用户表', '系统登录用户信息'),
        ('t_role', '角色表', '用户角色定义'),
        ('t_permission', '权限表', '系统权限清单'),
        ('t_role_permission', '角色权限关联表', '角色与权限的多对多关联'),
        ('t_archive_template', '归档模板表', 'Word归档模板文件管理'),
        ('t_archive_record', '归档记录表', '归档材料生成记录'),
        ('t_eval_device', '测评设备表', '硬件测评设备档案'),
        ('t_pentest_tool', '渗透工具表', '渗透测试软件工具信息'),
        ('t_dict', '字典主表', '数据字典分类'),
        ('t_dict_item', '字典项表', '数据字典具体项目'),
        ('t_sys_config', '系统配置表', '系统参数配置'),
        ('t_login_log', '登录日志表', '用户登录记录'),
        ('t_operation_log', '操作日志表', '系统操作记录'),
        ('t_user_preference', '用户偏好表', '用户个性化配置'),
    ]
    
    db_table = doc.add_table(rows=len(db_overview)+1, cols=3)
    set_table_borders(db_table)
    db_headers = ['表名', '中文名称', '说明']
    for j, h in enumerate(db_headers):
        db_table.rows[0].cells[j].text = h
        set_cell_background(db_table.rows[0].cells[j], '1F3864')
        for p in db_table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (a, b, c) in enumerate(db_overview):
        db_table.rows[i+1].cells[0].text = a
        db_table.rows[i+1].cells[1].text = b
        db_table.rows[i+1].cells[2].text = c
        if i % 2 == 0:
            for cell in db_table.rows[i+1].cells:
                set_cell_background(cell, 'F5F5F5')
    
    heading(doc, '4.2  核心数据表详细说明', 2)
    
    def add_table_def(doc, table_name, table_desc, columns):
        heading(doc, f'{table_name}  {table_desc}', 3)
        t = doc.add_table(rows=len(columns)+1, cols=5)
        set_table_borders(t)
        headers = ['字段名', '类型', '是否加密', '是否必填', '说明']
        for j, h in enumerate(headers):
            t.rows[0].cells[j].text = h
            set_cell_background(t.rows[0].cells[j], '2E75B6')
            for p in t.rows[0].cells[j].paragraphs:
                for r in p.runs:
                    r.font.color.rgb = RGBColor(255,255,255)
                    r.font.bold = True
                    r.font.size = Pt(10)
        for i, col in enumerate(columns):
            for j, val in enumerate(col):
                t.rows[i+1].cells[j].text = val
                for p in t.rows[i+1].cells[j].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10)
        doc.add_paragraph()
    
    # t_project
    add_table_def(doc, 't_project', '项目主表', [
        ('id', 'BIGINT', '否', '是', '主键，自增'),
        ('project_no', 'VARCHAR(50)', '否', '是', '项目编号，唯一约束'),
        ('project_name', 'VARCHAR(500)', '是(SM4)', '是', '项目名称'),
        ('customer_name', 'VARCHAR(500)', '是(SM4)', '是', '委托单位名称'),
        ('customer_address', 'VARCHAR(1000)', '是(SM4)', '否', '委托单位地址'),
        ('customer_contact', 'VARCHAR(500)', '是(SM4)', '否', '联系人姓名'),
        ('customer_phone', 'VARCHAR(500)', '是(SM4)', '否', '联系人电话'),
        ('system_name_merged', 'VARCHAR(1000)', '否', '否', '被测系统合并名称（缓存字段）'),
        ('sys_count_l2', 'INT', '否', '否', '二级系统数量'),
        ('sys_count_l3', 'INT', '否', '否', '三级系统数量'),
        ('project_type_id', 'BIGINT', '否', '否', '项目类型ID（关联t_dict_item）'),
        ('industry_id', 'BIGINT', '否', '否', '所属行业ID（关联t_dict_item）'),
        ('project_manager_id', 'BIGINT', '否', '否', '项目经理ID（关联t_staff）'),
        ('project_leader_id', 'BIGINT', '否', '否', '项目负责人ID（关联t_staff）'),
        ('contract_date', 'DATE', '否', '否', '合同签署日期'),
        ('contract_amount', 'VARCHAR(500)', '是(SM4)', '否', '合同金额'),
        ('paper_archived', 'TINYINT', '否', '否', '纸质归档状态：0否/1是'),
        ('electronic_archived', 'TINYINT', '否', '否', '电子归档状态：0否/1是'),
        ('phase_prepare', 'VARCHAR(100)', '否', '否', '准备阶段完成时间'),
        ('phase_plan', 'VARCHAR(100)', '否', '否', '计划阶段完成时间'),
        ('phase_onsite', 'VARCHAR(100)', '否', '否', '现场测评完成时间'),
        ('phase_report', 'VARCHAR(100)', '否', '否', '报告阶段完成时间'),
        ('task_appoint_date', 'DATE', '否', '否', '任务书签约时间'),
        ('year_belong', 'VARCHAR(10)', '否', '否', '所属年份'),
        ('business_person', 'VARCHAR(500)', '是(SM4)', '否', '业务人员'),
        ('project_region', 'VARCHAR(100)', '否', '否', '项目地区'),
        ('remark', 'TEXT', '是(SM4)', '否', '备注'),
        ('report_mail_date', 'DATE', '否', '否', '报告邮寄日期'),
        ('report_mail_no', 'VARCHAR(100)', '否', '否', '报告邮寄编号'),
        ('data_hmac', 'VARCHAR(200)', '否', '否', '数据完整性HMAC摘要'),
        ('created_at', 'DATETIME', '否', '是', '创建时间（自动填充）'),
        ('updated_at', 'DATETIME', '否', '是', '更新时间（自动填充）'),
    ])
    
    # t_project_system
    add_table_def(doc, 't_project_system', '被测系统表', [
        ('id', 'BIGINT', '否', '是', '主键，自增'),
        ('project_id', 'BIGINT', '否', '是', '所属项目ID（关联t_project）'),
        ('sys_seq', 'INT', '否', '是', '系统序号（1,2,3...）'),
        ('sys_name', 'VARCHAR(500)', '是(SM4)', '是', '被测系统名称'),
        ('sys_level', 'INT', '否', '是', '系统等级：2=二级，3=三级'),
        ('eval_index', 'VARCHAR(100)', '否', '否', '测评指标（如：MLPS-3）'),
        ('record_no', 'VARCHAR(100)', '否', '否', '备案编号'),
        ('data_hmac', 'VARCHAR(200)', '否', '否', '数据完整性HMAC摘要'),
        ('created_at', 'DATETIME', '否', '是', '创建时间'),
        ('updated_at', 'DATETIME', '否', '是', '更新时间'),
    ])
    
    # t_project_member
    add_table_def(doc, 't_project_member', '项目人员关联表', [
        ('id', 'BIGINT', '否', '是', '主键，自增'),
        ('project_id', 'BIGINT', '否', '是', '项目ID（关联t_project）'),
        ('member_id', 'BIGINT', '否', '是', '人员ID（关联t_staff）'),
        ('role_type', 'VARCHAR(50)', '否', '是', '角色类型，见下方角色列表'),
        ('data_hmac', 'VARCHAR(200)', '否', '否', '数据完整性HMAC摘要'),
        ('created_at', 'DATETIME', '否', '是', '创建时间'),
    ])
    
    add_para(doc, 'role_type 取值说明：')
    role_types = [
        ('project_manager', '项目经理'),
        ('project_leader', '项目负责人（自动同步自主表project_leader_id）'),
        ('registered_evaluator', '登记测评师（显示为项目组成员）'),
        ('actual_member', '实际参与人员'),
        ('survey_editor', '调研访谈编制人'),
        ('plan_editor', '测评方案编制人'),
        ('report_editor', '测评报告编制人'),
        ('network_evaluator', '网络测评人'),
        ('host_evaluator', '主机测评人'),
        ('physical_evaluator', '物理测评人'),
        ('tool_scanner', '工具扫描人'),
        ('pentest_member', '渗透测试人'),
    ]
    rt_table = doc.add_table(rows=len(role_types)+1, cols=2)
    set_table_borders(rt_table)
    rt_table.rows[0].cells[0].text = 'role_type值'
    rt_table.rows[0].cells[1].text = '说明'
    set_cell_background(rt_table.rows[0].cells[0], '4472C4')
    set_cell_background(rt_table.rows[0].cells[1], '4472C4')
    for c in rt_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(role_types):
        rt_table.rows[i+1].cells[0].text = k
        rt_table.rows[i+1].cells[1].text = v
    doc.add_paragraph()
    
    # t_staff
    add_table_def(doc, 't_staff', '人员清单表', [
        ('id', 'BIGINT', '否', '是', '主键，自增'),
        ('staff_no', 'VARCHAR(50)', '否', '否', '人员编号'),
        ('real_name', 'VARCHAR(500)', '是(SM4)', '是', '真实姓名'),
        ('department', 'VARCHAR(100)', '否', '否', '所属部门'),
        ('position', 'VARCHAR(100)', '否', '否', '职位'),
        ('role_level', 'VARCHAR(20)', '否', '否', '资质等级（高级/中级/初级）'),
        ('cert_no', 'VARCHAR(500)', '是(SM4)', '否', '证书编号'),
        ('cert_expire', 'DATE', '否', '否', '证书有效期'),
        ('phone', 'VARCHAR(500)', '是(SM4)', '否', '联系电话'),
        ('email', 'VARCHAR(500)', '是(SM4)', '否', '电子邮箱'),
        ('status', 'TINYINT', '否', '是', '状态：1=在职，0=离职'),
        ('data_hmac', 'VARCHAR(200)', '否', '否', '数据完整性HMAC摘要'),
        ('created_at', 'DATETIME', '否', '是', '创建时间'),
        ('updated_at', 'DATETIME', '否', '是', '更新时间'),
    ])
    
    # t_user
    add_table_def(doc, 't_user', '系统用户表', [
        ('id', 'BIGINT', '否', '是', '主键，自增'),
        ('username', 'VARCHAR(50)', '否', '是', '登录用户名，唯一'),
        ('real_name', 'VARCHAR(500)', '是(SM4)', '是', '真实姓名'),
        ('password_hash', 'VARCHAR(200)', '否', '是', 'BCrypt哈希密码（查询时不返回）'),
        ('role_id', 'BIGINT', '否', '是', '角色ID（关联t_role）'),
        ('phone', 'VARCHAR(500)', '是(SM4)', '否', '手机号'),
        ('email', 'VARCHAR(500)', '是(SM4)', '否', '邮箱'),
        ('totp_secret', 'VARCHAR(500)', '是(SM4)', '否', 'TOTP密钥'),
        ('totp_enabled', 'TINYINT', '否', '是', 'TOTP开关：0=关，1=开'),
        ('status', 'TINYINT', '否', '是', '状态：1=正常，0=停用'),
        ('login_fail_count', 'INT', '否', '否', '连续登录失败次数'),
        ('locked_until', 'DATETIME', '否', '否', '锁定至时间'),
        ('last_login_at', 'DATETIME', '否', '否', '最近登录时间'),
        ('last_login_ip', 'VARCHAR(50)', '否', '否', '最近登录IP'),
        ('password_changed_at', 'DATETIME', '否', '否', '密码修改时间'),
        ('first_login', 'TINYINT', '否', '否', '首次登录标志：1=是'),
        ('data_hmac', 'VARCHAR(200)', '否', '否', '数据完整性HMAC摘要'),
        ('created_at', 'DATETIME', '否', '是', '创建时间'),
        ('updated_at', 'DATETIME', '否', '是', '更新时间'),
    ])
    
    # t_archive_template
    add_table_def(doc, 't_archive_template', '归档模板表', [
        ('id', 'BIGINT', '否', '是', '主键，自增'),
        ('template_code', 'VARCHAR(100)', '否', '是', '模板编码（唯一）'),
        ('template_name', 'VARCHAR(200)', '否', '是', '模板名称'),
        ('template_category', 'VARCHAR(100)', '否', '否', '模板分类'),
        ('applicable_type', 'VARCHAR(100)', '否', '否', '适用项目类型'),
        ('applicable_level', 'INT', '否', '否', '适用系统等级（2/3）'),
        ('file_path', 'VARCHAR(500)', '否', '是', '文件存储相对路径'),
        ('file_original_name', 'VARCHAR(200)', '否', '是', '原始文件名'),
        ('file_size', 'BIGINT', '否', '否', '文件大小（字节）'),
        ('version', 'VARCHAR(20)', '否', '否', '模板版本号'),
        ('is_default', 'TINYINT', '否', '否', '是否默认模板：1=是'),
        ('status', 'TINYINT', '否', '是', '状态：1=启用，0=停用'),
        ('upload_by', 'BIGINT', '否', '否', '上传人用户ID'),
        ('upload_at', 'DATETIME', '否', '是', '上传时间'),
        ('data_hmac', 'VARCHAR(200)', '否', '否', '数据完整性HMAC摘要'),
    ])
    
    # t_eval_device
    add_table_def(doc, 't_eval_device', '测评设备表', [
        ('id', 'BIGINT', '否', '是', '主键，自增'),
        ('device_no', 'VARCHAR(50)', '否', '否', '设备编号'),
        ('device_name', 'VARCHAR(200)', '否', '是', '设备名称'),
        ('device_model', 'VARCHAR(200)', '否', '否', '设备型号'),
        ('device_type', 'INT', '否', '是', '设备类型：1=测评专用机，2=扫描设备'),
        ('owner_staff_id', 'BIGINT', '否', '否', '归属人员ID（关联t_staff）'),
        ('status', 'TINYINT', '否', '是', '状态：1=正常，0=停用'),
        ('remark', 'VARCHAR(500)', '否', '否', '备注'),
        ('data_hmac', 'VARCHAR(200)', '否', '否', '数据完整性HMAC摘要'),
    ])
    
    add_page_break(doc)
    
    # ===== CHAPTER 5: ENCRYPTED FIELDS =====
    heading(doc, '5  数据加密字段说明', 1)
    add_para(doc, '系统采用SM4（国密对称加密算法）对以下字段进行加密存储，加密密钥通过环境变量 APP_ROOT_KEY 配置，切勿泄露。')
    
    encrypted_fields = [
        ('t_project', 'project_name', '项目名称'),
        ('t_project', 'customer_name', '委托单位名称'),
        ('t_project', 'customer_address', '委托单位地址'),
        ('t_project', 'customer_contact', '联系人姓名'),
        ('t_project', 'customer_phone', '联系人电话'),
        ('t_project', 'contract_amount', '合同金额'),
        ('t_project', 'business_person', '业务人员'),
        ('t_project', 'remark', '备注'),
        ('t_project_system', 'sys_name', '被测系统名称'),
        ('t_staff', 'real_name', '人员真实姓名'),
        ('t_staff', 'cert_no', '证书编号'),
        ('t_staff', 'phone', '联系电话'),
        ('t_staff', 'email', '电子邮箱'),
        ('t_user', 'real_name', '用户真实姓名'),
        ('t_user', 'phone', '用户手机号'),
        ('t_user', 'email', '用户邮箱'),
        ('t_user', 'totp_secret', 'TOTP密钥'),
    ]
    
    enc_table = doc.add_table(rows=len(encrypted_fields)+1, cols=3)
    set_table_borders(enc_table)
    enc_headers = ['数据表', '字段名', '说明']
    for j, h in enumerate(enc_headers):
        enc_table.rows[0].cells[j].text = h
        set_cell_background(enc_table.rows[0].cells[j], '1F3864')
        for p in enc_table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (a, b, c) in enumerate(encrypted_fields):
        enc_table.rows[i+1].cells[0].text = a
        enc_table.rows[i+1].cells[1].text = b
        enc_table.rows[i+1].cells[2].text = c
        if i % 2 == 0:
            for cell in enc_table.rows[i+1].cells:
                set_cell_background(cell, 'FFF2CC')
    
    add_para(doc, '\n注意事项：')
    add_para(doc, '1. 在进行数据库迁移、备份恢复时，请确保加密密钥（APP_ROOT_KEY）保持一致，否则已加密数据将无法正确解密；', indent=0.5)
    add_para(doc, '2. 修改 APP_ROOT_KEY 后，所有已加密数据都需要重新加密，操作前请做好完整备份；', indent=0.5)
    add_para(doc, '3. HMAC摘要字段（data_hmac）用于数据完整性校验，不可手动修改数据库中的被保护字段，否则会导致HMAC校验失败。', indent=0.5)
    
    add_page_break(doc)
    
    # ===== CHAPTER 6: DEPLOYMENT STRUCTURE =====
    heading(doc, '6  系统部署结构', 1)
    
    add_para(doc, '系统采用前后端分离部署方式，推荐的生产环境部署结构如下：')
    
    deploy_info = [
        ('数据库服务器', '172.16.30.9', 'CentOS 8', 'MySQL 8.0', '数据库存储服务'),
        ('应用服务器', '172.16.40.8', 'CentOS 8', 'Nginx + Java 17 + Spring Boot', '后端API服务 + 前端静态文件服务'),
    ]
    
    dep_table = doc.add_table(rows=len(deploy_info)+1, cols=5)
    set_table_borders(dep_table)
    dep_headers = ['服务器角色', 'IP地址', '操作系统', '主要软件', '职责']
    for j, h in enumerate(dep_headers):
        dep_table.rows[0].cells[j].text = h
        set_cell_background(dep_table.rows[0].cells[j], '1F3864')
        for p in dep_table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, row_data in enumerate(deploy_info):
        for j, val in enumerate(row_data):
            dep_table.rows[i+1].cells[j].text = val
    
    doc.add_paragraph()
    add_para(doc, '网络访问流程：')
    flow_para = doc.add_paragraph()
    flow_run = flow_para.add_run('用户浏览器  →  HTTPS(443)  →  Nginx(172.16.40.8)  →  [静态文件 or 反向代理]  →  Spring Boot API(8080)  →  MySQL(172.16.30.9:3306)')
    flow_run.font.name = 'Courier New'
    flow_run.font.size = Pt(10)
    flow_para.paragraph_format.left_indent = Cm(1)
    
    doc.add_paragraph()
    add_para(doc, '部署路径规划（推荐）：')
    path_data = [
        ('/opt/djbh/', '应用根目录'),
        ('/opt/djbh/backend/', '后端JAR包目录'),
        ('/opt/djbh/frontend/', '前端静态文件目录'),
        ('/opt/djbh/data/templates/', 'Word模板存储目录'),
        ('/opt/djbh/data/archives/', '生成的归档材料目录'),
        ('/opt/djbh/logs/', '应用日志目录'),
        ('/opt/djbh/ssl/', 'SSL证书目录'),
    ]
    path_table = doc.add_table(rows=len(path_data)+1, cols=2)
    set_table_borders(path_table)
    path_table.rows[0].cells[0].text = '路径'
    path_table.rows[0].cells[1].text = '说明'
    set_cell_background(path_table.rows[0].cells[0], '2E75B6')
    set_cell_background(path_table.rows[0].cells[1], '2E75B6')
    for c in path_table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255,255,255)
                r.font.bold = True
    for i, (k, v) in enumerate(path_data):
        path_table.rows[i+1].cells[0].text = k
        path_table.rows[i+1].cells[1].text = v
    
    # Save
    output_path = os.path.join(OUTPUT_DIR, '文档1-系统功能及开发设计架构说明.docx')
    doc.save(output_path)
    print(f"Doc1 saved: {output_path}")
    print(f"File size: {os.path.getsize(output_path)} bytes")

create_doc1()
